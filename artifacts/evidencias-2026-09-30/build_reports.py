from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(__file__).resolve().parent
AUTHORS = 'Roberto Enrique Quezada-Rodríguez y Melanie Celeste Tello Fuentes'
PAPER_DIR = OUT.parent / 'manuscript-review-2026-09-29/rgn-2026-09-30'

def font(run, size=11, bold=False, color='000000'):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + key), 'Times New Roman')
    for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        rf.attrib.pop(qn('w:' + key), None)

def base(short):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin = s.bottom_margin = Inches(.7)
    s.left_margin = s.right_margin = Inches(.75)
    s.header_distance = s.footer_distance = Inches(.3)
    for name, size in [('Normal',11), ('Title',20), ('Heading 1',14), ('Heading 2',12)]:
        st = d.styles[name]
        st.font.name = 'Times New Roman'
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0,0,0)
        st.paragraph_format.space_after = Pt(7)
        st.paragraph_format.line_spacing = 1.08
        if name != 'Normal':
            st.paragraph_format.keep_with_next = True
        rf = st.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            rf.set(qn('w:'+key), 'Times New Roman')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            rf.attrib.pop(qn('w:'+key), None)
        for b in list(st.element.xpath('./w:pPr/w:pBdr')):
            b.getparent().remove(b)
    h = s.header.paragraphs[0]
    font(h.add_run('QUEZADA RODRÍGUEZ · TELLO FUENTES   |   '+short),9,color='555555')
    f = s.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(f.add_run('30 de septiembre de 2026   ·   Página '),9,color='555555')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
    d.core_properties.author = AUTHORS
    d.core_properties.subject = short
    return d

def p(d, text, style=None, size=11):
    x = d.add_paragraph(style=style)
    font(x.add_run(text), size, bool(style and style.startswith('Heading')))
    return x

def title(d, text, subtitle):
    p(d,text,'Title',20)
    p(d,subtitle,size=12)
    p(d,AUTHORS)
    p(d,'Universidad Nacional de Trujillo · Perú')
    p(d,'Fecha de preparación y consulta de fuentes: 30 de septiembre de 2026',size=10)

def head(d,text): p(d,text,'Heading 1',14)
def sub(d,text): p(d,text,'Heading 2',12)
def page(d): d.add_page_break()

def table(d, headers, rows, widths):
    t = d.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for c,w in zip(t.columns,widths): c.width = Inches(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'C8C8C8'); borders.append(e)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','bottom','left','right']:
        e=OxmlElement('w:'+side); e.set(qn('w:w'),'85'); e.set(qn('w:type'),'dxa'); margins.append(e)
    props.append(margins)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for n, values in enumerate([headers]+rows):
        row=t.rows[0] if n==0 else t.add_row()
        no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for c,text,w in zip(row.cells,values,widths):
            c.width=Inches(w)
            q=c.paragraphs[0]; q.paragraph_format.space_after=Pt(2); q.paragraph_format.line_spacing=1.02
            font(q.add_run(text),10.5,n==0)
            if n==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EDEDED'); c._tc.get_or_add_tcPr().append(shade)
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def link(d,label,url):
    q=d.add_paragraph(); q.paragraph_format.space_after=Pt(4)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),q.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    rf=OxmlElement('w:rFonts'); rf.set(qn('w:ascii'),'Times New Roman'); rf.set(qn('w:hAnsi'),'Times New Roman'); pr.append(rf)
    co=OxmlElement('w:color'); co.set(qn('w:val'),'174F78'); pr.append(co)
    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'20'); pr.append(sz)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); pr.append(u)
    r.append(pr); tx=OxmlElement('w:t'); tx.text=label; r.append(tx); h.append(r); q._p.append(h)

def save(d, filename):
    for q in d.paragraphs:
        q.paragraph_format.widow_control=True
        for b in list(q._p.xpath('./w:pPr/w:pBdr')): b.getparent().remove(b)
    d.save(OUT/filename)

def practice():
    d=base('Informe técnico de la práctica')
    title(d,'Informe técnico del proyecto sslope','Entregable 1 · Evidencia técnica para la práctica académica')
    head(d,'Objetivo y alcance')
    p(d,'Desarrollamos un prototipo para visualizar taludes y evaluar pronósticos de desplazamiento mediante el método de elementos finitos (FEM) y redes de memoria a corto y largo plazo (LSTM). Este informe documenta el código y los artefactos disponibles, y explica cómo sustentan el artículo. No certifica asistencia, una actividad ejecutada en aula ni una validación en mina; esas evidencias no fueron proporcionadas.')
    p(d,'El artículo estudia el caso semisintético TA-01. Los módulos de Century Mine y el laboratorio de entrenamiento son desarrollos independientes y no deben presentarse como resultados de campo que validen ese artículo.')
    head(d,'Autoría y contribuciones')
    table(d,['Integrante','Contribución declarada'],[
        ['Roberto Enrique Quezada-Rodríguez','Software, análisis y redacción. Autor de correspondencia: requezadaro@unitru.edu.pe.'],
        ['Melanie Celeste Tello Fuentes','Análisis, redacción, revisión y validación. ORCID: 0009-0007-4728-138X.'],
    ],[2.75,4.25])
    head(d,'Tres productos de la entrega')
    table(d,['Producto','Estado y contenido'],[
        ['Informe de práctica','Este documento reúne arquitectura, trazabilidad, procedimiento de demostración y límites.'],
        ['Artículo de investigación','Manuscrito_TA01_RGN.docx ya existe. Está adaptado a Rudarsko-geološko-naftni zbornik, con tablas y figuras. No es una constancia de envío o aceptación.'],
        ['Reporte de revistas','Documento separado con cuatro opciones, requisitos de revisión, periodicidad, cuartiles, costos y fuentes.'],
    ],[1.7,5.3])
    p(d,'No se ha generado un archivo comprimido. Los apellidos del equipo figuran en los nombres de los dos nuevos informes.',size=10)

    page(d); head(d,'Componentes y archivos del proyecto')
    p(d,'La aplicación principal utiliza Node.js, JavaScript y Three.js. El servidor conserva telemetría en SQLite y expone una API local. El visor muestra geometría, lluvia y desplazamientos; la amplificación visual facilita observar deformaciones pequeñas, pero no cambia las unidades ni las métricas.')
    table(d,['Componente','Archivos o carpeta','Función'],[
        ['Visor y API','public/\nsrc/server.js\nsrc/persistence.js','Interfaz 3D, controles, consultas y almacenamiento.'],
        ['Mecánica TA-01','src/core/fem-2d.js\nsrc/core/study-case-ta01.js','FEM triangular lineal 2D y definición del caso supuesto.'],
        ['Datos y particiones','scripts/generate-fem-dataset.js\nscripts/split-fem-chronological.js','Generación de escenarios y separación temporal.'],
        ['Aprendizaje','scripts/train-lstm.py\nscripts/train-physics-guided.py','Entrenamiento temporal y corrector residual guiado por física.'],
        ['Evaluación','data/validation/\nsrc/core/research-engine.js','Comparaciones, incertidumbre y verificaciones archivadas.'],
        ['Century Mine','docs/CENTURY_DATASET.md\nsrc/core/century-dataset.js','Preparación de observaciones públicas y evaluación retrospectiva independiente.'],
        ['Laboratorio separado','research-app/','Análisis exploratorio de datos (EDA), limpieza, ajuste y comparación. Versión Next.js/FastAPI y entrega posterior Python/Streamlit.'],
    ],[1.25,2.8,2.95])
    sub(d,'Distinción entre las aplicaciones')
    p(d,'El visor del artículo funciona en el puerto 3000. El laboratorio anterior usa interfaz 3001 y API 8001; la entrega Streamlit documentada utiliza 8501. Sus datos, historial y resultados son independientes. La presencia del laboratorio no cambia automáticamente el protocolo ni las conclusiones del manuscrito TA-01.')

    page(d); head(d,'Cómo se obtiene la evidencia del artículo')
    p(d,'TA-01 combina lluvia diaria de NASA POWER/MERRA-2 para 2020–2025 con una geometría y propiedades geotécnicas supuestas. La lluvia es reanálisis, no una estación instalada en la mina. El desplazamiento objetivo lo genera el FEM; por eso el conjunto es semisintético.')
    table(d,['Etapa','Evidencia conservada','Qué permite comprobar'],[
        ['Escenarios','data/generated/ta01-fem-500-manifest.json','Configuración e identificación de los 500 escenarios.'],
        ['Partición inicial','ta01-replication-manifest.json','350/75/75 escenarios, agrupados por fecha de lluvia. No es una prueba cronológica.'],
        ['Prueba 2024','Manifiesto de replicación','258/82/82 escenarios de entrenamiento, validación y prueba; 78 escenarios futuros excluidos.'],
        ['Prueba 2025','Manifiesto de replicación','340/82/78 escenarios de entrenamiento, validación y prueba.'],
        ['Comparación','ta01-rolling-origin-model-comparison.csv','MAE, RMSE y R² del pronóstico a 1 h y 6 h.'],
        ['Incertidumbre','ta01-rolling-origin-bootstrap.json','5000 remuestreos agrupados por fecha e intervalos del 95 %. No es bootstrap de bloques consecutivos.'],
    ],[1.3,2.7,3.0])
    head(d,'Preguntas y conclusiones trazables')
    p(d,'RQ1 compara persistencia, ridge, LSTM e híbrido. El ranking cambia entre las pruebas cronológicas. RQ2 evalúa el corrector: LSTM obtiene mejores errores en 2024 y el híbrido en 2025; los cuatro intervalos pareados incluyen cero. No hay evidencia de ventaja consistente ni una prueba de equivalencia.')
    p(d,'RQ3 muestra deterioro ante ruido y datos faltantes, con mayor degradación del híbrido en las condiciones evaluadas. RQ4 verifica internamente el FEM y una aproximación espacial bajo el mismo operador. Compartir el solver de referencia no constituye validación geotécnica independiente.')
    p(d,'MAE es el error absoluto medio; RMSE, la raíz del error cuadrático medio; R², el coeficiente de determinación. No se añaden valores p: el manuscrito no informa pruebas de hipótesis cuya ejecución reproducible no esté conservada. Las cifras detalladas permanecen en sus tablas y archivos originales.')

    page(d); head(d,'Verificación de esta entrega')
    table(d,['Comprobación','Resultado al 30 de septiembre de 2026'],[
        ['Integridad de artefactos TA-01','86 de 86 archivos enumerados en sources coinciden con el SHA-256 de ta01-replication-manifest.json. Sin ausencias ni discrepancias en ese conjunto.'],
        ['Pruebas automatizadas','Node.js v24.19.0: 117 pruebas aprobadas, ninguna fallida, en 25 archivos. La primera ejecución restringida bloqueó un servidor local y la extracción del ZIP; la repetición autorizada sin esas restricciones resolvió ambos bloqueos. No es una medida de exactitud científica.'],
        ['Alcance de la revisión','No se reentrenaron modelos ni se regeneraron resultados científicos. Los hashes comprueban integridad, no reproducibilidad bit a bit del entrenamiento.'],
    ],[1.75,5.25])
    head(d,'Procedimiento para demostrar el software')
    for text in [
        '1. Desde la raíz del proyecto, ejecutar npm start con Node.js 24 o superior y las dependencias instaladas. Abrir http://127.0.0.1:3000.',
        '2. Seleccionar el caso TA-01. Mostrar el FEM 2D y las comparaciones a 1 h y 6 h; contrastar los valores con los artefactos de data/validation.',
        '3. Activar la deformación del visor y alternar 1× con una escala ampliada. Explicar que la escala modifica la representación, no el desplazamiento calculado.',
        '4. Mostrar Century Mine por separado. Identificar sus observaciones como históricas y las predicciones como evaluación retrospectiva, no monitoreo actual.',
        '5. Para EDA y ajuste independiente, usar research-app/ y su README. No mezclar ese historial con los experimentos TA-01 del artículo.',
    ]: p(d,text)
    sub(d,'Límites y siguiente validación')
    p(d,'El prototipo permite comparar métodos, estudiar sensibilidad y comprobar partes de la implementación. No permite certificar estabilidad de una mina ni emitir alertas operacionales. Para ello se requieren geometría y materiales del sitio, condiciones hidráulicas, sensores calibrados y evaluación prospectiva independiente. Los datos públicos de Century ayudan a estudiar pronósticos observados, pero no aportan todos los parámetros necesarios para validar el FEM del sitio.')
    sub(d,'Disponibilidad y uso de IA')
    p(d,'Fuentes locales: README.md; docs/CASO_TA01_FEM.md; docs/CENTURY_DATASET.md; research-app/README.md; data/validation/ta01-replication-manifest.json. Artículo existente: artifacts/manuscript-review-2026-09-29/rgn-2026-09-30/Manuscrito_TA01_RGN.docx.',size=10)
    p(d,'OpenAI Codex se utilizó como asistencia para desarrollo de software, revisión documental y preparación editorial. Los autores son responsables de verificar el contenido, interpretar los resultados y aprobar la versión que presenten. Financiación y conflictos de interés: ninguno declarado por los autores.',size=10)
    save(d,'Quezada_Rodriguez_Tello_Fuentes_01_Practica_tecnica.docx')

def journals():
    d=base('Reporte de revistas')
    title(d,'Reporte de revistas de acceso abierto','Entregable 3 · Opciones Q2 y Q3 sin APC para el manuscrito TA-01')
    head(d,'Resultado de la búsqueda')
    p(d,'Recomendamos mantener Rudarsko-geološko-naftni zbornik como primera opción por su alcance en minería y geotecnia, ausencia de tasas para autores y requisitos ya atendidos en el manuscrito. El reporte identifica cuatro alternativas o candidatas; su inclusión no implica aceptación probable ni compatibilidad completa con la versión actual.')
    p(d,'Se verificaron políticas editoriales públicas. Los cuartiles se indican con sistema, año y categoría: SJR utiliza datos de Scopus y no es el mismo indicador que CiteScore; JCR/JIF corresponde a Clarivate. Estar indexada no equivale a pertenecer a Q2 o Q3. La clasificación cambia con el año y el área.')
    table(d,['Revista','Cuartil documentado','Números por año','Revisores a proponer'],[
        ['RGN · Croacia','Q3 SJR y JCR 2025; Q2 CiteScore 2025 en dos áreas','5','2 obligatorios'],
        ['Geologia Croatica · Croacia','Q3 JCR/JIF 2024, Geology; actualizar antes del envío','3','3 o 5 recomendados'],
        ['Geofizika · Croacia','Q3 JCR/JIF 2025 en dos áreas','2','3–5 recomendados'],
        ['CAMES · Polonia','Q3 SJR 2025 según fuente secundaria; confirmación directa pendiente','4','3–5 obligatorios'],
    ],[1.6,2.65,.85,1.9])
    p(d,'Todas ofrecen publicación ordinaria sin APC según sus políticas. Geologia Croatica y Geofizika contemplan costos opcionales para color en papel. Se puede evitar esa opción y mantener color gratuito en línea. APC significa cargo por procesamiento del artículo.',size=10)
    sub(d,'Cómo leer las fichas')
    p(d,'“Propuestos” son candidatos que identifican los autores; “asignados” son evaluadores que elige el editor. Una recomendación de candidatos no es una obligación, y ninguno de ellos ha sido contactado en esta búsqueda. Los plazos descritos por las revistas no garantizan una decisión ni una fecha de publicación.')
    p(d,'El país de la revista no es un criterio de facilidad de aceptación. La prioridad depende del alcance, del rigor del trabajo y de los requisitos verificables. No debe enviarse el mismo manuscrito a varias revistas a la vez.',size=10)

    page(d); head(d,'Rudarsko geološko naftni zbornik')
    p(d,'Nombre oficial: Rudarsko-geološko-naftni zbornik (RGN). Editor: Universidad de Zagreb, Facultad de Minería, Geología e Ingeniería del Petróleo, Croacia. ISSN 0353-4529; eISSN 1849-0409.')
    table(d,['Aspecto','Información verificada'],[
        ['Indexación y cuartil','Scopus y WoS. Para 2025: SJR 0.332, Q3 en cinco categorías; JIF 1.3, Q3 en Geosciences, Multidisciplinary y Mining & Mineral Processing. [R1]'],
        ['Distinción Q2','CiteScore 2.8: Q2 en Geology y General Earth and Planetary Sciences; Q3 en Geotechnical Engineering and Engineering Geology, General Energy y Water Science and Technology. No confundir con SJR. [R1]'],
        ['Costos y frecuencia','Sin tasas para autores de cualquier país; acceso abierto. Cinco números regulares al año desde 2021, además de posibles especiales. El apartado histórico de cuatro números está desactualizado. [R2–R3]'],
        ['Revisión por pares','Al menos dos evaluadores independientes; revisión simple ciego, con modalidad abierta opcional. [R2]'],
        ['Candidatos propuestos','Dos obligatorios en carta al editor: nombres, títulos, áreas, afiliaciones y correos. El editor puede no utilizarlos. [R3]'],
        ['Preparación','Plantilla oficial Word, máximo 25 páginas formateadas; resumen inglés y croata de hasta 250 palabras. Carta con originalidad y revisión profesional o nativa del inglés británico. [R3]'],
        ['Plazos publicados','Primera ronda de revisión y corrección: tres semanas cada una; siguientes: una semana. El editor aspira a procesar trabajos regulares en unos tres meses si se cumplen los plazos; no es una garantía. [R2–R3]'],
    ],[1.45,5.55])
    sub(d,'Encaje y estado de preparación')
    p(d,'Encaje alto, como evaluación computacional semisintética de taludes. El manuscrito y la carta con dos revisores ya existen. Quedan aprobación final de ambos autores, confirmación de ausencia de conflictos con los candidatos y revisión competente del inglés y del resumen croata. Debe mantenerse la declaración real de uso de IA y evitar describir el prototipo como gemelo operacional validado.')
    sub(d,'Fuentes verificables')
    link(d,'R1 · Anuncio editorial del 21 de septiembre de 2026 con métricas 2025','https://ojs.srce.hr/rgn/announcement')
    link(d,'R2 · Política, costos, frecuencia y revisión por pares','https://ojs.srce.hr/rgn/about')
    link(d,'R3 · Instrucciones de envío y carta al editor','https://ojs.srce.hr/rgn/about/submissions')

    page(d); head(d,'Geologia Croatica')
    p(d,'Editada por el Croatian Geological Survey y la Croatian Geological Society, Croacia. ISSN 1330-030X; eISSN 1333-4875. Incluye geología de ingeniería y acepta autores sin restricciones de nacionalidad. [G1]')
    table(d,['Aspecto','Información verificada'],[
        ['Indexación y cuartil','Scopus y WoS/SCIE. Q3 JCR/JIF 2024, Geology; JIF 1.1. La tabla institucional corresponde a la edición publicada en 2025, no a métricas 2025. No se afirma aquí un Q2 actualizado sin comprobación directa. [G2–G3]'],
        ['Costos y frecuencia','Tres números al año. Sin tasas ordinarias de publicación. Color en línea sin recargo; color en la edición impresa es opcional y puede cobrarse, con importe a consultar. [G1, G4]'],
        ['Revisión por pares','Dos revisores; Hrčak describe evaluación internacional con anonimato simple. [G1]'],
        ['Candidatos propuestos','La guía invita a proponer 3 o 5, con afiliación, dirección y correo. Recomendación, no requisito expresado como obligatorio. Permite indicar hasta tres no preferidos por conflictos. [G4]'],
        ['Idioma y formato','Inglés británico. Word, letra de 12 puntos, doble espacio y márgenes de 2.5 cm. Resumen de 50–300 palabras; figuras y tablas como archivos separados. [G4]'],
        ['Envío y plazos','Guía: solicitar acceso a editorial.office@geologia-croatica.hr. Plazo típico declarado desde primer envío hasta publicación: 3–9 meses, variable. Revisión lingüística posterior por su editor en Reino Unido. [G4]'],
    ],[1.45,5.55])
    sub(d,'Encaje y reservas')
    p(d,'Encaje medio. Para un caso idealizado fuera de sus regiones de interés, el aporte a geología de ingeniería debe quedar claro y no limitarse al software. El cuartil más reciente necesita comprobarse en JCR antes de considerarla definitivamente elegible para la exigencia del docente.')
    p(d,'La página editorial consultada contiene enlaces de pie de página ajenos a la revista. No se utilizaron como fuentes. Antes de ingresar credenciales o cargar archivos, conviene confirmar la ruta de envío con el contacto institucional publicado en Hrčak. Este hallazgo no permite por sí solo calificar a la revista como ilegítima.',size=10)
    sub(d,'Fuentes verificables')
    link(d,'G1 · Hrčak institucional, periodicidad, revisión y ausencia de tasas','https://hrcak.srce.hr/en/geologia-croatica')
    link(d,'G2 · Indexación y política editorial','https://www.geologia-croatica.hr/index.php/GC/about')
    link(d,'G3 · Tabla institucional JCR con año de datos 2024','https://vvg.hr/znanost-i-istrazivanje/2025/10/16/journal-citation-reports-2025/')
    link(d,'G4 · Guía de envío, revisores, formato y costos opcionales','https://www.geologia-croatica.hr/index.php/GC/about/submissions')

    page(d); head(d,'Geofizika')
    p(d,'Editada por el Instituto Geofísico Andrija Mohorovičić, Departamento de Geofísica, Facultad de Ciencias, Universidad de Zagreb, Croacia. ISSN 0352-3659; eISSN 1846-6346. [F1]')
    table(d,['Aspecto','Información verificada'],[
        ['Indexación y cuartil','Scopus y WoS/SCIE. JIF 2025 de 1.5 y Q3 en Geochemistry & Geophysics y Geosciences, Multidisciplinary, según la página editorial vigente. [F1–F2]'],
        ['Costos y frecuencia','Dos números al año; calendario declarado 31 de enero y 31 de julio, con artículos in press. Sin APC ordinario; color gratuito en línea. [F1, F3]'],
        ['Excepción de costos','Color en papel opcional: una página EUR 150; dos, EUR 120 por página; tres, EUR 100 por página; cuatro o más, EUR 80 por página. No elegir esta opción si se requiere costo cero. [F3]'],
        ['Revisión por pares','Dos evaluadores para artículos originales, comunicaciones cortas y revisiones; uno para comunicaciones preliminares y trabajos profesionales. [F1, F3]'],
        ['Candidatos propuestos','De 3 a 5 recomendados, sin relación personal o profesional actual con los autores. La selección definitiva corresponde al editor. [F3]'],
        ['Idioma y preparación','Inglés preferido o croata. Resumen hasta 250 palabras y traducción croata; el editor puede ayudar con esta última. Declaraciones de datos y autoría obligatorias. [F3]'],
        ['Política de IA','Permite mejorar legibilidad de texto propio, con revisión humana; prohíbe crear texto nuevo o alterar imágenes y videos mediante IA. Su compatibilidad con el uso real de Codex requiere aclaración editorial. [F3]'],
    ],[1.45,5.55])
    sub(d,'Encaje y reservas')
    p(d,'Encaje bajo a medio: prioriza problemas geofísicos, incluida hidrología, no cualquier estudio geotécnico. El efecto de la lluvia en TA-01 es semisintético y no debe convertirse en una afirmación de calibración hidrológica. No recomendamos envío inmediato de la versión actual por el alcance y la restricción sobre redacción asistida por IA. No hay un plazo total de decisión garantizado en las páginas consultadas.')
    sub(d,'Fuentes verificables')
    link(d,'F1 · Información general, indexación y frecuencia','https://ojs.srce.hr/geofizika/about')
    link(d,'F2 · Página editorial con JIF y cuartiles 2025','https://ojs.srce.hr/geofizika/index')
    link(d,'F3 · Revisores, formato, excepción de costos y política de IA','https://ojs.srce.hr/geofizika/about/submissions')

    page(d); head(d,'Computer Assisted Methods in Engineering and Science')
    p(d,'CAMES. Editor: Institute of Fundamental Technological Research, Polish Academy of Sciences, Polonia. ISSN 2299-3649; eISSN 2956-5839. Orientada a modelado y simulación computacional en ingeniería. [C1, C2]')
    table(d,['Aspecto','Información verificada'],[
        ['Indexación y cuartil','Indexación Scopus confirmada por el editor. SJR 2025 de 0.298 confirmado en su web; Q3 en Computational Mechanics, Computer Science Applications y Mechanical Engineering según fuente secundaria. [C1–C2]'],
        ['Grado de comprobación','El cuartil requiere verificación directa en SCImago o Scopus antes del envío. No se atribuye un cuartil JCR ni indexación WoS no acreditados en estas fuentes. [C1–C2]'],
        ['Costos y frecuencia','Sin costo de publicación. Trimestral: cuatro números al año, además de publicación Online First. [C1, C3]'],
        ['Revisión y candidatos','Al menos dos evaluadores independientes, simple ciego. Carta con 3–5 candidatos y justificación, externos a la institución; la guía pide considerar candidatos de otros países. [C3–C4]'],
        ['Idioma y formato','Inglés revisado por hablante nativo. Envío inicial PDF, doble espacio, 12 puntos y máximo 40000 caracteres tipográficos, aproximadamente 20 páginas. [C3]'],
        ['Límite ambiguo','La guía dice literalmente 200 typographic signs para el resumen. No se interpreta como 200 palabras: debe consultarse al editor antes de adaptar el manuscrito. [C3]'],
        ['Plazos e IA','Los revisores disponen de unas cuatro semanas; no equivale al plazo total. La declaración de envío exige no haber usado IA para escribir el artículo, con excepciones para gramática y usos de investigación. Es una barrera para la versión actual asistida por Codex. [C1, C4–C5]'],
    ],[1.45,5.55])
    sub(d,'Encaje y decisión')
    p(d,'Encaje temático alto para métodos computacionales, pero opción condicionada. Requeriría demostrar novedad metodológica y aclarar por escrito el uso real de IA antes de enviar. No corresponde marcar una declaración de no uso de IA si no es cierta. Por esa razón no es una alternativa inmediata al manuscrito ya preparado para RGN.')
    sub(d,'Fuentes verificables')
    link(d,'C1 · Editor, alcance, indexación, métricas y periodicidad','https://cames.ippt.pan.pl/index.php/cames/about')
    link(d,'C2 · Cuartiles 2025 en fuente secundaria, pendientes de confirmar directamente','https://researchjournalrank.com/journal/computer-assisted-methods-in-engineering-and-science')
    link(d,'C3 · Guía, costos y revisores sugeridos','https://cames.ippt.pan.pl/index.php/cames/author-guidelines')
    link(d,'C4 · Procedimiento de revisión por pares','https://cames.ippt.pan.pl/index.php/cames/author-guidelines/peer-review-process')
    link(d,'C5 · Ética y declaración de uso de IA','https://cames.ippt.pan.pl/index.php/cames/journal-policies')

    page(d); head(d,'Decisión y requisitos antes del envío')
    p(d,'Primera opción: RGN. Ya disponemos de adaptación editorial y carta con revisores. Segunda opción de estudio: Geologia Croatica, tras actualizar cuartil y consultar encaje. Geofizika y CAMES se conservan como opciones condicionadas, principalmente por alcance y políticas de IA. No se asigna una probabilidad de aceptación sin evidencia.')
    table(d,['Requisito','Situación o acción necesaria'],[
        ['Manuscrito y carta RGN','Existen Manuscrito_TA01_RGN.docx y Carta_editor_RGN_con_revisores.docx en artifacts/manuscript-review-2026-09-29/rgn-2026-09-30/.'],
        ['Revisores RGN','Dos candidatos documentados: Jinsong Huang, University of Newcastle, Australia; Gianvito Scaringi, Charles University, República Checa. Los datos de contacto y fundamentos están en la carta. No fueron contactados.'],
        ['Conflictos y autoría','Ambos autores deben aprobar el envío y confirmar ausencia de vínculos que comprometan imparcialidad de los candidatos. La selección final es editorial.'],
        ['Idioma','No declarar una revisión nativa o profesional del inglés que aún no se haya realizado. Revisar también el borrador del resumen croata.'],
        ['Cuartil exigido','Acordar con el docente si acepta SJR, CiteScore o JCR, y qué año requiere. RGN documenta Q3 SJR/JCR 2025 y Q2 CiteScore 2025 en áreas concretas.'],
        ['Integridad científica','Mantener datos semisintéticos, resultados sin ventaja consistente del corrector, límites de campo y declaración de Codex. No esconder asistencia de IA.'],
    ],[1.55,5.45])
    head(d,'Opciones descartadas del filtro de costo cero')
    p(d,'Acceso abierto no significa necesariamente publicación gratuita para autores. Los siguientes ejemplos no cumplen el filtro económico o no tienen clasificación actual suficientemente favorable para esta entrega.')
    link(d,'Studia Geotechnica et Mechanica · APC de EUR 450 más IVA desde 2025','https://sgem.pwr.edu.pl/index.php?id=7')
    link(d,'Acta Geodynamica et Geomaterialia · Tasa ordinaria de EUR 400; reducida para estudiantes','https://www2.irsm.cas.cz/index.php?page=acta_scope')
    link(d,'Acta Polytechnica · Candidata gratuita excluida de la lista principal por registro secundario Q4 en 2025','https://researchjournalrank.com/journal/acta-polytechnica')
    sub(d,'Límites de este reporte')
    p(d,'Es una consulta de fuentes públicas al 30 de septiembre de 2026, no una certificación emitida por Scopus o Clarivate. Los enlaces permiten revisar políticas. Geologia Croatica tiene evidencia de cuartil con año anterior; CAMES conserva una comprobación secundaria explícita. La falta de acceso directo a una base no se transforma en una confirmación. No se realizó un envío ni se obtuvo aceptación editorial.')
    p(d,'OpenAI Codex asistió en la búsqueda, síntesis y preparación de este reporte. Los autores deben revisar las políticas vigentes al presentar el artículo y asumir responsabilidad por sus declaraciones.',size=10)
    save(d,'Quezada_Rodriguez_Tello_Fuentes_03_Reporte_revistas_Q2_Q3.docx')

practice()
journals()
for file in sorted(OUT.glob('*.docx')):
    print(file)
