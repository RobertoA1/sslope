"""Populate the official BGM text template with the verified TA-01 manuscript."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Cm, Inches, Pt
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
TEMPLATE = Path("/home/robertoa1/Descargas/Manuscrito.docx")
EXPECTED_SHA256 = "950e4a16772d21e5d1f2868dd1170abad5c1734f19122bd4dc68de3593972338"
SOURCE = ROOT / "Articulo_TA01_corregido_BGM.docx"
OUTPUT = ROOT / "Manuscrito_TA01_BGM_plantilla.docx"

assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() == EXPECTED_SHA256
base = Document(TEMPLATE)
source = Document(SOURCE)
sp = list(source.paragraphs)

body = base._element.body
for node in list(body):
    if node.tag != qn("w:sectPr"):
        body.remove(node)


def add(text="", role="body", *, bold=None, italic=None):
    styles = {
        "body": "Acknowledgment Head",
        "title": "Art title",
        "author": "Author",
        "affiliation": "Author Affiliation",
        "abstract_head": "Abstract Head",
        "abstract": "Abstract",
        "head1": "Head 1",
        "head2": "Acknowledgment Head",
        "reference_head": "Reference Head",
        "reference": "References",
        "caption": "References",
    }
    para = base.add_paragraph(style=styles[role])
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.line_spacing = 2.0
    if role in {"body", "abstract"}:
        para.paragraph_format.first_line_indent = Cm(0.5 if role == "body" else 0)
    else:
        para.paragraph_format.first_line_indent = Cm(0)
    if role == "reference":
        para.paragraph_format.left_indent = Cm(0.75)
        para.paragraph_format.first_line_indent = -Cm(0.75)
    if text:
        r = para.add_run(text)
        if bold is not None:
            r.bold = bold
        if italic is not None:
            r.italic = italic
        if role == "body":
            r.bold = False
    return para


def copy_text(par, role="body"):
    out = add(role=role)
    for run in par.runs:
        if not run.text:
            continue
        text = run.text.replace(" & ", " and ")
        new = out.add_run(text)
        new.bold = run.bold if run.bold is not None else (True if role in {"head1", "head2"} else False)
        new.italic = run.italic
        new.underline = run.underline
    if not out.text and par.text:
        out.add_run(par.text.replace(" & ", " and "))
    return out


def keypoints(label, lines):
    add(label, "abstract", bold=True)
    for line in lines:
        add(line.removeprefix("• "), "abstract")


# Template front matter: English first, then Spanish, exactly as the sample.
add(sp[1].text.strip(), "title")
es_title = add(sp[0].text.strip(), "title")
for run in es_title.runs:
    run.font.size = Pt(14)
add("Título abreviado: Pronóstico semisintético FEM–LSTM en taludes", "abstract", italic=True)
add("Roberto Enrique Quezada-Rodríguez¹*, Melanie Celeste Tello Fuentes¹", "author")
add(
    "¹Universidad Nacional de Trujillo, Ciudad Universitaria, Av. Juan Pablo II S/N, "
    "Urb. San Andrés, Trujillo, La Libertad, Perú",
    "affiliation",
)
add("*Autor de correspondencia: requezadaro@unitru.edu.pe (Roberto Enrique Quezada-Rodríguez)", "abstract")
add("ORCID:", "abstract_head")
add("Roberto Enrique Quezada-Rodríguez: 0009-0004-0869-3881", "abstract")
add("Melanie Celeste Tello Fuentes: 0009-0007-4728-138X", "abstract")
add("ABSTRACT", "abstract_head")
english_abstract = copy_text(sp[11], "abstract")
for run in english_abstract.runs:
    if "5.000" in run.text:
        run.text = run.text.replace("5.000", "5000")
copy_text(sp[12], "abstract")
keypoints("Key points:", [sp[i].text for i in (18, 19, 20)])
add("RESUMEN", "abstract_head")
copy_text(sp[8], "abstract")
copy_text(sp[9], "abstract")
keypoints("Puntos clave:", [sp[i].text for i in (14, 15, 16)])


def add_equation(n, text=""):
    para = add(role="body")
    para.paragraph_format.first_line_indent = Cm(0)
    para.paragraph_format.tab_stops.add_tab_stop(Inches(2.95), WD_TAB_ALIGNMENT.CENTER)
    para.paragraph_format.tab_stops.add_tab_stop(Inches(5.85), WD_TAB_ALIGNMENT.RIGHT)
    para.add_run("\t")
    if n == 1:
        para.add_run().add_picture(str(ROOT / "image2.png"), width=Inches(4.2))
    elif n == 4:
        math = para.add_run("∂cφ/∂r ≥ 0,    ∂cφ/∂s ≤ 0")
        math.font.name = "Cambria Math"
        math.font.size = Pt(11)
    else:
        math = para.add_run(text.split("\t")[0])
        math.font.name = "Cambria Math"
        math.font.size = Pt(11)
    para.add_run(f"\t({n})")


skip_caption = re.compile(r"^(?:Figura|Figure|Tabla|Table) [1-6](?:\.|\n|$)")
for i in range(23, 158):
    current = sp[i]
    text = current.text.strip()
    if not text:
        continue
    if skip_caption.match(text) or text.startswith("Nota."):
        continue
    if i in {58, 60, 61, 62}:
        add_equation({58: 1, 60: 2, 61: 3, 62: 4}[i], text)
        continue
    if i == 59:
        add(text + " En (4), r es la lluvia futura acumulada y s el índice de seguridad simulado; las derivadas se refieren solo a la salida del corrector cφ, con las demás entradas constantes.")
        continue
    if current.style.name == "Heading 1":
        copy_text(current, "head1")
    elif current.style.name == "Heading 2":
        copy_text(current, "head2")
    elif text.startswith("Financiación."):
        add("Financiación. Los autores no recibieron financiación, becas ni apoyo institucional específico para este estudio.")
    elif text.startswith("Conflictos de interés."):
        add("Conflictos de interés. Los autores declaran que no tienen conflictos de interés relacionados con este trabajo.")
    elif text.startswith("Uso de inteligencia artificial generativa."):
        add(
            "Uso de inteligencia artificial. OpenAI Codex se utilizó como asistente en el desarrollo "
            "del software y los scripts del prototipo, el cotejo del manuscrito con el código y los "
            "archivos de validación, la preparación técnica de figuras y tablas a partir de resultados "
            "archivados, y la revisión editorial y maquetación de esta versión. Su uso no sustituye "
            "la responsabilidad de los autores sobre los métodos, las cifras, las referencias, la "
            "interpretación y la versión presentada."
        )
    else:
        copy_text(current)

# Acknowledgements are optional in the journal template. The authors have not
# supplied names or institutions to thank, so none are invented here.
add("Referencias", "reference_head")
for par in sp[159:205]:
    if par.text.strip():
        copy_text(par, "reference")


figures = {
    1: ("Flujo de datos y validación del prototipo TA-01. LHS: muestreo por hipercubo latino; FEM: elementos finitos; LSTM: memoria a corto y largo plazo. Esquema semisintético, no sistema de alerta de campo.",
        "Data flow and validation of the TA-01 prototype. LHS: Latin hypercube sampling; FEM: finite element method; LSTM: long short-term memory. Semisynthetic scheme, not a field alert system."),
    2: ("Sensibilidad del desplazamiento inducido a la discretización FEM. Se muestran número de nodos, desplazamiento máximo y diferencia frente a la malla 48 × 32; es una referencia interna, no convergencia demostrada. Caso semisintético. Fuente: data/validation/ta01-fem-mesh-summary.csv.",
        "FEM mesh sensitivity of rainfall-induced displacement. Node count, maximum displacement, and difference relative to the 48 × 32 mesh are shown; this is an internal reference, not demonstrated convergence. Semisynthetic case. Source: data/validation/ta01-fem-mesh-summary.csv."),
    3: ("Diferencia pareada de MAE entre LSTM e híbrido. Puntos: Δ = MAE(LSTM) − MAE(híbrido); barras: intervalos percentiles del 95 % agrupados por fecha. Las escalas horizontales difieren; no se representan valores p. Caso semisintético. Fuente: data/validation/ta01-rolling-origin-bootstrap.json.",
        "Paired MAE difference between LSTM and the hybrid. Points show Δ = MAE(LSTM) − MAE(hybrid); bars are date-clustered 95% percentile intervals. Horizontal scales differ; no p values are shown. Semisynthetic case. Source: data/validation/ta01-rolling-origin-bootstrap.json."),
    4: ("Cobertura retrospectiva de los intervalos del híbrido. Todas las ventanas frente a lluvia ≥5 mm/día; la línea discontinua indica el nivel nominal. El grupo de lluvia alta comprende dos fechas en 2024 y tres en 2025. Las ventanas están correlacionadas. Caso semisintético. Fuente: data/validation/ta01-rolling-interval-coverage.csv.",
        "Retrospective coverage of hybrid prediction intervals. All windows are compared with rainfall ≥5 mm/day; the dashed line is nominal coverage. The high-rainfall group contains two dates in 2024 and three in 2025. Windows are correlated. Semisynthetic case. Source: data/validation/ta01-rolling-interval-coverage.csv."),
    5: ("Sensibilidad predictiva a ruido y faltantes aleatorios. Ruido relativo a la desviación de entrenamiento y faltantes imputados con medias de entrenamiento. Cada punto compara la condición alterada con el caso limpio del mismo modelo; no hubo reentrenamiento. Caso semisintético. Fuente: data/validation/ta01-robustness-summary.csv.",
        "Predictive sensitivity to noise and random missing inputs. Noise is relative to the training standard deviation, and missing values are imputed with training means. Each point compares a perturbed condition with the clean case for the same model; no retraining was performed. Semisynthetic case. Source: data/validation/ta01-robustness-summary.csv."),
    6: ("Enmascaramiento de entradas hidrológicas y mecánicas. Las entradas se sustituyeron por medias de entrenamiento; MAE en las mismas ventanas del conjunto no cronológico. Las variantes no se reentrenaron ni aíslan términos de la pérdida. Caso semisintético. Fuente: data/validation/ta01-ablation-summary.csv.",
        "Masking of hydrological and mechanical inputs. Inputs were replaced by training means; MAE was evaluated on the same non-chronological test windows. The variants were not retrained and do not isolate individual loss terms. Semisynthetic case. Source: data/validation/ta01-ablation-summary.csv."),
}
tables = {
    1: ("Comparación del alcance con trabajos directamente relacionados. Síntesis de las obras citadas; no constituye una revisión sistemática.",
        "Scope comparison with closely related studies. Synthesis of the cited studies; this is not a systematic review."),
    2: ("Configuración documentada del experimento. Arquitectura y entrenamiento según data/models/ta01-lstm-1h.json y ta01-physics-guided-1h.json; las corridas de 6 h y cronológicas registran las épocas terminadas. Python y hardware no están fijados.",
        "Documented experimental configuration. Architecture and training follow data/models/ta01-lstm-1h.json and ta01-physics-guided-1h.json; 6-h and chronological runs record completed epochs. Python version and hardware are not fixed."),
    3: ("Error absoluto medio de la prueba separada por fecha de lluvia. MAE en mm; las cuatro alternativas se evaluaron en las mismas ventanas de cada horizonte. Fuente: data/validation/ta01-model-comparison.csv.",
        "Mean absolute error in the rainfall-date-separated test. MAE in mm; all four methods were evaluated on the same windows at each horizon. Source: data/validation/ta01-model-comparison.csv."),
    4: ("MAE de LSTM e híbrido en pruebas cronológicas. MAE en mm. Fuentes: data/validation/ta01-rolling-origin-model-comparison.csv y ta01-rolling-origin-bootstrap.json.",
        "LSTM and hybrid MAE in chronological tests. MAE in mm. Sources: data/validation/ta01-rolling-origin-model-comparison.csv and ta01-rolling-origin-bootstrap.json."),
    5: ("Alcance de la verificación numérica y espacial. Valores contrastados con data/validation/ta01-fem-mesh-sensitivity.json y ta01-spatial-pinn-validation.json; no son errores de pronóstico temporal.",
        "Scope of numerical and spatial verification. Values checked against data/validation/ta01-fem-mesh-sensitivity.json and ta01-spatial-pinn-validation.json; these are not temporal forecast errors."),
    6: ("MAE con entradas limpias y perturbadas. Valores archivados sin recalcular; las respuestas son semisintéticas y no hubo reentrenamiento bajo perturbación. Fuente: data/validation/ta01-robustness-summary.csv.",
        "MAE with clean and perturbed inputs. Archived values were not recalculated; responses are semisynthetic and no retraining was performed under perturbation. Source: data/validation/ta01-robustness-summary.csv."),
}

add("Pies de figuras y tablas / Figure and table captions", "head1")
for n, (es, en) in figures.items():
    add(f"Figura {n}. {es}", "caption")
    add(f"Figure {n}. {en}", "caption")
for n, (es, en) in tables.items():
    add(f"Tabla {n}. {es}", "caption")
    add(f"Table {n}. {en}", "caption")

base.save(OUTPUT)
print(OUTPUT)
