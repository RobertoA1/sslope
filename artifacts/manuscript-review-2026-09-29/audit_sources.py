from pathlib import Path
from docx import Document
from zipfile import ZipFile
from PIL import Image
from lxml import etree
import json, re, hashlib, urllib.request, urllib.parse, concurrent.futures, csv, io

ROOT=Path('/home/robertoa1/Datos/proyectos/sslope')
OUT=Path(__file__).parent
SOURCE=Path('/home/robertoa1/Descargas/Artículo - Prototipo semisintético FEM LSTM con corrección guiada por física para pronosticar desplazamientos en taludes.docx')
d=Document(SOURCE)
OUT.mkdir(parents=True,exist_ok=True)
records=[]
for i,p in enumerate(d.paragraphs):
    records.append({'index':i,'style':p.style.name,'text':p.text})
(OUT/'original-extracted.json').write_text(json.dumps({'paragraphs':records,'tables':[[[c.text for c in r.cells] for r in t.rows] for t in d.tables]},ensure_ascii=False,indent=2))
refs=[r for r in records if 109<=r['index']<=149]
def get(url):
    req=urllib.request.Request(url,headers={'User-Agent':'Scientific-manuscript-reference-audit/1.0','Accept':'application/json'})
    with urllib.request.urlopen(req,timeout=28) as f:return json.load(f)
def verify(r):
    m=re.search(r'https://doi.org/(\S+)',r['text'])
    if not m:return {**r,'status':'website','doi':None}
    doi=m.group(1); info={**r,'doi':doi}
    try:
        j=get('https://api.crossref.org/works/'+urllib.parse.quote(doi,safe=''))['message']
        info.update(status='registered',registry='Crossref',metadata=j)
    except Exception as e:
        try:
            j=get('https://api.datacite.org/dois/'+urllib.parse.quote(doi,safe=''))['data']['attributes']
            info.update(status='registered',registry='DataCite',metadata=j)
        except Exception as e2:info.update(status='VERIFY',error=str(e),fallback_error=str(e2))
    return info
with concurrent.futures.ThreadPoolExecutor(max_workers=6) as pool: checked=list(pool.map(verify,refs))
(OUT/'references-audit.json').write_text(json.dumps(checked,ensure_ascii=False,indent=2))
for r in checked:
    m=r.get('metadata',{}); print(r['index'],r['status'],r['doi'],m.get('title',m.get('titles')),m.get('published',m.get('publicationYear')))

mf=json.loads((ROOT/'data/validation/ta01-replication-manifest.json').read_text())
hashes=[]
for name,record in mf['sources'].items():
    p=ROOT/name
    actual=hashlib.sha256(p.read_bytes()).hexdigest() if p.exists() else None
    hashes.append({'file':name,'expected':record['sha256'],'actual':actual,'match':actual==record['sha256']})
(OUT/'hash-audit.json').write_text(json.dumps(hashes,indent=2))
print('HASH AUDIT',len(hashes),'mismatch',[r['file'] for r in hashes if not r['match']])
images=[]
with ZipFile(SOURCE) as z:
    for n in z.namelist():
        if 'word/media/' not in n:continue
        content=z.read(n); p=OUT/Path(n).name; p.write_bytes(content)
        im=Image.open(io.BytesIO(content)); images.append({'file':n,'pixels':im.size,'metadata':im.info})
    root=etree.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    print('IMAGES',images)
    print('DRAWINGS',len(root.findall('.//w:drawing',ns)),'VML',len(root.findall('.//w:pict',ns)))
(OUT/'images-audit.json').write_text(json.dumps(images,default=str,ensure_ascii=False,indent=2))
