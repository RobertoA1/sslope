"""Adapt the archived TA-01 manuscript to the retained official RGN template."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import csv
import hashlib
import json
import re

import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from PIL import Image
from docx import Document
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from rgn_content import *

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[1]
OUT = ROOT / 'rgn-2026-09-30'
FIG = OUT / 'figures'
FIG.mkdir(parents=True, exist_ok=True)
TEMPLATE = Path('/tmp/RGN_official_template.docx')
TEMPLATE_HASH = 'bb10cc69c58e6f7394bdc0b69a90ce435f77c2dab0e0c99520000c93db34967c'
assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() == TEMPLATE_HASH
reference = Document(TEMPLATE)
source = Document(ROOT / 'Articulo_TA01_corregido_BGM.docx')
doc = Document(TEMPLATE)
body = doc.element.body
for element in list(body):
    if element.tag != qn('w:sectPr'):
        body.remove(element)

# Do not redefine template styles; clone the template's real paragraph patterns.
PATTERNS = {'body':14, 'title':0, 'author':2, 'affiliation':3, 'abstract_heading':7,
            'abstract':8, 'keywords':10, 'h1':12, 'h2':25, 'caption':31,
            'reference':66, 'smallheading':83}

def p(text='', role='body', bold=False, italic=False):
    xml = deepcopy(reference.paragraphs[PATTERNS[role]]._p)
    par = Paragraph(xml, doc._body)
    par.clear()
    body.insert(len(body)-1, xml)
    fmt = par.paragraph_format
    fmt.line_spacing = 1.0
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.keep_with_next = role in {'h1','h2','smallheading','abstract_heading'}
    fmt.keep_together = role in {'title','author','caption'}
    fmt.widow_control = True
    fmt.left_indent = Cm(0)
    fmt.right_indent = Cm(0)
    fmt.first_line_indent = Cm(.63 if role in {'body','abstract'} else 0)
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if role == 'reference':
        fmt.left_indent = Cm(.63)
        fmt.first_line_indent = Cm(-.63)
    if role in {'caption'}:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if role in {'h1','h2','smallheading','abstract_heading','title','author','affiliation','keywords'}:
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if role in {'h1','h2','smallheading'}:
        fmt.space_before = Pt(10)
        fmt.space_after = Pt(10)
    if text:
        run = par.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14 if role == 'title' else 12 if role in {'h1','smallheading'} else 9 if role == 'affiliation' else 10)
        run.font.color.rgb = RGBColor(0,0,0)
        run.bold = bold or role in {'title','author','h1','h2','smallheading','abstract_heading'}
        run.italic = italic
        lang = OxmlElement('w:lang'); lang.set(qn('w:val'),'en-GB'); run._r.get_or_add_rPr().append(lang)
    return par

def hyperlink(par, label, target, size=9):
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'),doc.part.relate_to(target,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman');pr.append(fonts)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(size*2));pr.append(sz)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);par._p.append(h)

def caption(label, text, *, before=False):
    par=p('', 'caption')
    par.paragraph_format.space_before=Pt(10)
    par.paragraph_format.space_after=Pt(10)
    par.paragraph_format.keep_with_next=before
    for txt,b in [(label+'. ',True),(text,False)]:
        r=par.add_run(txt);r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=b
    return par

TABLE_WIDTHS = [[2.8,3.7,3.7,2.9,3.9],[3.0,8.8,5.2],[2.0,2.2,3.2,3.2,3.2,3.2],
                [3.0,2.5,2.7,4.4,4.4],[4.0,3.2,3.0,6.8],[2.1,3.7,2.5,4.3,4.4]]

def wrap_path(text):
    return re.sub(r'\b(?:data|scripts|docs)/\S+', lambda m:m[0].replace('/','/\u200b').replace('-','-\u200b'), text)

def table(rows, widths):
    t=doc.add_table(rows=len(rows),cols=len(rows[0]))
    t._tbl.remove(t._tbl.tblPr)
    t._tbl.insert(0,deepcopy(reference.tables[0]._tbl.tblPr))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    tw=t._tbl.tblPr.find(qn('w:tblW'))
    tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(round(sum(widths)*567)))
    for column,width in zip(t.columns,widths):column.width=Cm(width)
    for i,row in enumerate(t.rows):
        if i==0:
            header=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(header)
        nosplit=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(nosplit)
        for j,(cell,value) in enumerate(zip(row.cells,rows[i])):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil')
                x.set(qn('w:sz'),'4');x.set(qn('w:color'),'000000');borders.append(x)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge,amount in [('top',55),('bottom',55),('left',75),('right',75)]:
                x=OxmlElement('w:'+edge);x.set(qn('w:w'),str(amount));x.set(qn('w:type'),'dxa');margins.append(x)
            pr.append(margins)
            cp=cell.paragraphs[0];cp.style=doc.styles['Text']
            cp.paragraph_format.space_before=Pt(0);cp.paragraph_format.space_after=Pt(0)
            cp.paragraph_format.first_line_indent=Cm(0);cp.paragraph_format.line_spacing=1.0
            cp.paragraph_format.left_indent=Cm(0);cp.paragraph_format.right_indent=Cm(0)
            cp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or any(c.isalpha() for c in value) else WD_ALIGN_PARAGRAPH.CENTER
            if i and re.fullmatch(r'[16] h',value):cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=cp.add_run(wrap_path(value));r.font.name='Times New Roman';r.font.size=Pt(9);r.bold=i==0
    p('').paragraph_format.space_after=Pt(0)
    return t

def add_table(n):
    caption(f'Table {n}',TABLE_CAPTIONS[n-1],before=True)
    result=table(TABLES[n-1],TABLE_WIDTHS[n-1])
    if n==4:
        # Keep the small chronological comparison intact near page boundaries.
        for row in result.rows[:-1]:
            for cell in row.cells:
                cell.paragraphs[0].paragraph_format.keep_with_next=True
    return result

# Editable Word math, using the simple OMML structures already supported by
# the official template. Their final rendering is checked page by page.
def mr(text):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r

def ms(base,sub):
    s=OxmlElement('m:sSub');e=OxmlElement('m:e');su=OxmlElement('m:sub')
    e.append(mr(base) if isinstance(base,str) else base);su.append(mr(sub));s.append(e);s.append(su);return s

def hat(base):
    a=OxmlElement('m:acc');pr=OxmlElement('m:accPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),'̂');pr.append(ch)
    e=OxmlElement('m:e');e.append(mr(base));a.append(pr);a.append(e);return a

def frac(top,bottom):
    f=OxmlElement('m:f');nu=OxmlElement('m:num');de=OxmlElement('m:den')
    for x in top:nu.append(x)
    for x in bottom:de.append(x)
    f.append(nu);f.append(de);return f

def equation(n):
    par=p('');par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.first_line_indent=Cm(0)
    par.paragraph_format.space_before=Pt(10);par.paragraph_format.space_after=Pt(10)
    par.paragraph_format.tab_stops.add_tab_stop(Cm(8.1),WD_TAB_ALIGNMENT.CENTER)
    par.paragraph_format.tab_stops.add_tab_stop(Cm(17),WD_TAB_ALIGNMENT.RIGHT)
    par.add_run('\t');m=OxmlElement('m:oMath')
    if n==1:
        parts=[mr('K(θ)'),ms('u','t'),mr(' = '),ms('f','g'),mr('(θ) + '),ms('f','p'),mr('('),ms('p','t'),mr(', θ)')]
    elif n in [2,3]:
        parts=[ms(hat('y'),'LSTM' if n==2 else 'H'),mr('(t+h) = max{'),ms('y','t'),mr(', '),ms('y','t'),mr(' + '),ms('μ','Δ'),mr(' + '),ms('σ','Δ')]
        if n==3:parts.append(mr('['))
        parts.extend([ms('q','ψ'),mr('('),ms('X','t,h'),mr(')')])
        if n==3:parts.extend([mr(' + '),ms('c','φ'),mr('('),ms('Z','t,h'),mr(')]')])
        parts.append(mr('}'))
    else:
        parts=[frac([mr('∂'),ms('c','φ')],[mr('∂r')]),mr(' ≥ 0,    '),frac([mr('∂'),ms('c','φ')],[mr('∂s')]),mr(' ≤ 0')]
    for x in parts:m.append(x)
    par._p.append(m);r=par.add_run(f'\t({n})');r.font.name='Times New Roman';r.font.size=Pt(10)

def readcsv(name):
    with (REPO/name).open(newline='') as handle:return list(csv.DictReader(handle))
def jread(name):return json.loads((REPO/name).read_text())

plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.titlesize':11,'axes.labelsize':10,
                     'axes.spines.top':False,'axes.spines.right':False,'figure.facecolor':'white','savefig.facecolor':'white'})
BLUE='#0072B2';ORANGE='#D55E00';GRAY='#555555';PROVENANCE={}
def savefig(fig,n,files):
    file=FIG/f'Figure_{n}.jpg'
    fig.savefig(file,dpi=400,bbox_inches='tight',pad_inches=.15,pil_kwargs={'quality':95});plt.close(fig)
    PROVENANCE[str(n)]={'source_files':files,'sha256':{x:hashlib.sha256((REPO/x).read_bytes()).hexdigest() for x in files},
                        'classification':'illustrative workflow' if n==1 else 'archived semisynthetic results'}

f,ax=plt.subplots(figsize=(6.7,3));ax.axis('off');ax.set_xlim(0,3);ax.set_ylim(0,2.5)
labels=[('External rainfall','NASA POWER daily'),('Assumed scenarios','Geometry and LHS\nmaterial properties'),('FEM 2D','24 quasi-static states'),
        ('Temporal windows','6 h history\n1 h / 6 h horizons'),('LSTM and correction','Monotonic residual\ncorrection'),('Evaluation','Disjoint dates\n2024 / 2025')]
positions=[(.04,1.55),(1.05,1.55),(2.06,1.55),(2.06,.4),(1.05,.4),(.04,.4)]
for (title,bodytext),(x,y) in zip(labels,positions):
    ax.add_patch(Rectangle((x,y),.90,.63,fill=False,edgecolor=GRAY,lw=1.1))
    ax.text(x+.45,y+.43,title,ha='center',va='center',fontsize=8.5,fontweight='bold')
    ax.text(x+.45,y+.16,bodytext,ha='center',va='center',fontsize=7.7)
for (x,y),(xx,yy) in zip(positions,positions[1:]):
    if y==yy:ax.annotate('',xy=(xx if xx>x else xx+.9,yy+.31),xytext=(x+.9 if xx>x else x,y+.31),arrowprops={'arrowstyle':'->','color':GRAY})
    else:ax.annotate('',xy=(xx+.45,yy+.63),xytext=(x+.45,y),arrowprops={'arrowstyle':'->','color':GRAY})
ax.text(1.5,.06,'TA-01 semisynthetic workflow; not field monitoring.',ha='center',fontsize=8)
savefig(f,1,['scripts/generate-fem-dataset.js','scripts/train-lstm.py','scripts/train-physics-guided.py','scripts/generate-research-summary.js'])

mesh=readcsv('data/validation/ta01-fem-mesh-summary.csv')
f,axs=plt.subplots(1,2,figsize=(6.7,2.9),layout='constrained');x=[int(r['node_count']) for r in mesh]
axs[0].plot(x,[float(r['maximum_rainfall_displacement_mm']) for r in mesh],'o-',color=BLUE)
axs[0].set_ylabel('Maximum rainfall-induced\ndisplacement (mm)');axs[0].set_xlabel('Mesh nodes');axs[0].set_title('a) Computed magnitude')
axs[1].plot(x,[float(r['displacement_difference_against_48x32_percent']) for r in mesh],'s-',color=ORANGE)
axs[1].set_ylabel('Difference from 48 × 32 (%)');axs[1].set_xlabel('Mesh nodes');axs[1].set_title('b) Internal reference')
for ax in axs:ax.grid(alpha=.2);ax.tick_params(labelsize=8)
savefig(f,2,['data/validation/ta01-fem-mesh-summary.csv'])

boot=jread('data/validation/ta01-rolling-origin-bootstrap.json')['comparisons']
f,axs=plt.subplots(1,2,figsize=(6.7,2.9),layout='constrained')
for ax,h in zip(axs,[1,6]):
    rows=[r for r in boot if r['horizonHours']==h];scale=1e-7 if h==1 else 1e-6
    for yi,r in enumerate(rows):
        lo,hi=r['confidenceInterval95Mm'];v=r['observedDifferenceMm']
        ax.errorbar(v/scale,yi,xerr=[[(v-lo)/scale],[(hi-v)/scale]],fmt='o',color=BLUE if r['testYear']==2024 else ORANGE,capsize=4)
    ax.axvline(0,color=GRAY,ls='--',lw=1);ax.set_yticks([0,1],['2024 (82 dates)','2025 (76 dates)'])
    ax.set_xlabel(f'Δ MAE (× {"10⁻⁷" if h==1 else "10⁻⁶"} mm)');ax.set_title(f'{h} h: Δ > 0 favours hybrid')
    ax.grid(axis='x',alpha=.2);ax.tick_params(labelsize=8)
savefig(f,3,['data/validation/ta01-rolling-origin-bootstrap.json'])

cov=readcsv('data/validation/ta01-rolling-interval-coverage.csv')
f,ax=plt.subplots(figsize=(6.7,2.9),layout='constrained');keys=[(2024,1),(2024,6),(2025,1),(2025,6)];x=np.arange(4)
for off,segment,label,color,hatch in [(-.17,'ALL','All windows',BLUE,''),(.17,'RAIN_HIGH_GE_5','Rainfall ≥5 mm/day',ORANGE,'//')]:
    vals=[100*float(next(r for r in cov if int(r['test_year'])==y and int(r['horizon_hours'])==h and r['segment']==segment)['empirical_coverage']) for y,h in keys]
    ax.bar(x+off,vals,width=.33,color=color,hatch=hatch,label=label)
ax.axhline(95,color=GRAY,ls='--',label='Nominal 95%');ax.set_xticks(x,[f'{y} · {h} h' for y,h in keys]);ax.set_ylabel('Empirical coverage (%)')
ax.set_ylim(0,113);ax.legend(fontsize=8,loc='upper left',ncol=2);ax.grid(axis='y',alpha=.2)
savefig(f,4,['data/validation/ta01-rolling-interval-coverage.csv'])

rob=readcsv('data/validation/ta01-robustness-summary.csv');conditions=['CLEAN','NOISE_1PCT_STD','NOISE_5PCT_STD','MISSING_10PCT','MISSING_30PCT']
labels=['Clean','1% noise','5% noise','10% missing','30% missing']
f,axs=plt.subplots(1,2,figsize=(6.7,3.1),layout='constrained')
for ax,h in zip(axs,[1,6]):
    for model,c,marker in [('LSTM',BLUE,'o'),('HYBRID',ORANGE,'s')]:
        vals=[float(next(r for r in rob if int(r['horizon_hours'])==h and r['condition']==condition and r['model']==model)['mae_increase_percent']) for condition in conditions]
        ax.plot(range(5),vals,marker+'-',color=c,label='Hybrid' if model=='HYBRID' else model)
    ax.set_xticks(range(5),labels,rotation=35,ha='right',fontsize=8)
    ax.set_ylabel('MAE increase from clean (%)');ax.set_title(f'{h} h');ax.grid(alpha=.2);ax.legend(fontsize=8)
savefig(f,5,['data/validation/ta01-robustness-summary.csv'])

abl=readcsv('data/validation/ta01-ablation-summary.csv');variants=['LSTM','HYBRID_NO_HYDROLOGY','HYBRID_NO_FEM','HYBRID_FULL']
labels=['LSTM','Hydrology\nmask','Mechanics\nmask','Full\nhybrid']
f,axs=plt.subplots(1,2,figsize=(6.7,2.9),layout='constrained')
for ax,h in zip(axs,[1,6]):
    scale=1e-5 if h==1 else 1e-4
    vals=[float(next(r for r in abl if int(r['horizon_hours'])==h and r['variant_id']==v)['mae_mm'])/scale for v in variants]
    ax.bar(range(4),vals,color=[BLUE,'#888888','#BBBBBB',ORANGE],edgecolor=GRAY)
    ax.set_xticks(range(4),labels,fontsize=8);ax.set_ylabel(f'MAE (× {"10⁻⁵" if h==1 else "10⁻⁴"} mm)')
    ax.set_title(f'{h} h · no retraining');ax.grid(axis='y',alpha=.2)
savefig(f,6,['data/validation/ta01-ablation-summary.csv'])
(OUT/'figure-provenance.json').write_text(json.dumps(PROVENANCE,indent=2))

def add_figure(n):
    par=p('');par.paragraph_format.first_line_indent=Cm(0)
    par.paragraph_format.space_before=Pt(10);par.paragraph_format.keep_with_next=True
    par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(FIG/f'Figure_{n}.jpg'),width=Cm(17))
    caption(f'Figure {n}',FIGURE_CAPTIONS[n-1])

# First-page metadata follows the reserved title/author slots at the left of
# the template's floating editorial block.
title=p('Semisynthetic evaluation of\nFEM LSTM slope displacement\nforecasts','title')
title.paragraph_format.right_indent=Cm(7)
title.paragraph_format.space_after=Pt(10)
authors=p('Roberto Enrique Quezada-Rodríguez¹*\nMelanie Celeste Tello Fuentes¹','author')
authors.paragraph_format.right_indent=Cm(7)
authors.paragraph_format.space_after=Pt(10)
p('¹Universidad Nacional de Trujillo, Ciudad Universitaria, Av. Juan Pablo II S/N, Urb. San Andrés, Trujillo, La Libertad, Peru','affiliation')
ids=p('','affiliation');hyperlink(ids,'Roberto ORCID: 0009-0004-0869-3881','https://orcid.org/0009-0004-0869-3881')
ids.add_run('    ');hyperlink(ids,'Melanie ORCID: 0009-0007-4728-138X','https://orcid.org/0009-0007-4728-138X')
ids.paragraph_format.space_after=Pt(10)
p('Abstract','abstract_heading');p(ABSTRACT,'abstract')
p('Keywords: '+KEYWORDS,'keywords').paragraph_format.space_before=Pt(10)

for index,text in TEXT.items():
    if re.match(r'^\d+\.\s',text):
        p(text,'h1')
        if index==40:
            p('The methods connect the mechanical generator to temporal forecasts and their evaluation. The following subsections describe data provenance, model coupling, partitioning and reproducibility.')
        if index==71:
            p('Results are organised by research question. Numerical implementation checks are reported separately from temporal prediction and input robustness.')
    elif re.match(r'^\d+\.\d+\.',text):p(text,'h2')
    else:p(text)
    if index==32:add_table(1)
    if index==42:add_figure(1)
    if index==50:equation(1)
    if index==52:
        for n in [2,3,4]:equation(n)
    if index==62:add_table(2)
    if index==73:add_figure(2)
    if index==76:add_table(3)
    if index==79:add_table(4);add_figure(3)
    if index==82:add_figure(4)
    if index==85:add_table(5)
    if index==87:add_table(6);add_figure(5)
    if index==89:add_figure(6)

p('Data and code availability','smallheading');p(DATA_AVAILABILITY)
p('6. References','h1')
for old in source.paragraphs[159:205]:
    new=p('','reference')
    for oldrun in old.runs:
        if oldrun.text:
            txt=oldrun.text.replace('Publicación anticipada en línea.','Advance online publication.')
            if 'Dong, X.,' in old.text:
                txt=txt.replace('Xiang, H., Zhu, J., & Qiu, Y.','Xiang, H., … Qiu, Y.')
            if 'National Aeronautics and Space Administration' in old.text:
                txt=txt.replace('(s. f.)','(n.d.)').replace('[Conjunto de datos]','[Dataset]').replace('Recuperado el 22 de septiembre de 2026, de','Retrieved September 22, 2026, from')
            r=new.add_run(txt);r.font.name='Times New Roman';r.font.size=Pt(10);r.italic=oldrun.italic
    if not new.text:new.add_run(old.text.replace(' & ',' & '))

p('Appendix A Contextual literature map','h1')
p('Table A1 lists the forty contextual works retained from the manuscript bibliography. Topics are based on their research subjects; the inventory is not a systematic review, and their reported performance is not used as numerical evidence for TA-01. The directly related subset is analysed in Table 1.')
matrix=[['ID','Study','Methodological topic']]
for row,topic in zip(source.tables[6].rows[1:],LITERATURE_TOPICS):
    author=row.cells[1].text
    match=re.search(r'\((\d{4})\)',author);year=match[1]
    name=author.split(',',1)[0]
    if row.cells[0].text=='E-06':study='Zhang, Z., et al. (2024)'
    elif row.cells[0].text=='E-09':study='Zhang, S., et al. (2024)'
    elif row.cells[0].text=='E-01':study='Liu, X., et al. (2022)'
    elif row.cells[0].text=='E-02':study='Liu and Jia (2026)'
    else:study=name+' et al. ('+year+')'
    # Cite two- and single-author works correctly rather than using et al.
    replacements={'E-07':'Guan and Wang (2024)','E-11':'Yiğit and Şenol (2025)','E-12':'Dahal and Lombardo (2025)',
                  'E-17':'Nanda and Patel (2026)','S-02':'Bustillos Cotrado (2024)','S-09':'Blas Cano and Milla Vergara (2023)',
                  'S-12':'Montúfar Chiriboga (2025)','S-13':'Chilon Chilon (2025)','S-15':'Sequeira-Arguedas (2023)',
                  'S-16':'González Serrano (2025)','S-17':'Arias-Valencia and Aristizábal-Arias (2024)',
                  'S-19':'Zamora Nevado and Medina Vásquez (2023)'}
    study=replacements.get(row.cells[0].text,study)
    matrix.append([row.cells[0].text,study,topic])
caption('Table A1','Contextual methodological topics and their cited sources.',before=True)
table(matrix,[1.2,5.3,10.5])

p('SAŽETAK','smallheading');p(CROATIAN_TITLE,'h2');p(CROATIAN_ABSTRACT,'abstract')
p('Ključne riječi: '+CROATIAN_KEYWORDS,'keywords')
p('Authors’ contributions','smallheading');p(CONTRIBUTIONS)
p('Acknowledgements','smallheading');p(AI_DECLARATION)
p('Funding','smallheading');p('The authors received no funding, grants or specific institutional financial support for this study.')
p('Conflicts of interest','smallheading');p('The authors declare no conflicts of interest related to this work.')
p('Ethics','smallheading');p('The study uses public reanalysis rainfall and simulated geotechnical data. It involves no human participants, animals or personal information.')

# Fill correspondence and author placeholders, preserve publication fields as
# journal-assigned. Never claim an acceptance date or an assigned article DOI.
for par in doc.sections[0].header.paragraphs:
    for r in par.runs:
        if 'Surname' in r.text:r.text='R. E. Quezada-Rodríguez and M. C. Tello Fuentes'
footer=doc.sections[0].first_page_footer
for par in footer.paragraphs:
    text=par.text
    if 'Corresponding author' in text:
        par.clear();par.add_run('*Corresponding author: Roberto Enrique Quezada-Rodríguez')
    elif 'e-mail address' in text:
        par.clear();par.add_run('e-mail address: requezadaro@unitru.edu.pe')
    elif 'Received:' in text:
        par.clear();par.add_run('Received / Accepted / Available online: to be completed by editors')
    for r in par.runs:r.font.name='Times New Roman';r.font.size=Pt(9)
for par in doc.sections[0].footer.paragraphs:
    if '10.17794/rgn.xxxx' in par.text:
        par.clear();r=par.add_run('Rudarsko-geološko-naftni zbornik — manuscript for review');r.italic=True;r.font.name='Times New Roman';r.font.size=Pt(9)
for t in doc.sections[0].first_page_header._element.iter(qn('w:t')):
    if 'DOI: 10.17794/' in (t.text or ''):t.text='DOI: to be assigned by editors'
    if 'paper category' in (t.text or ''):t.text='Paper category: to be assigned by editors'

settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.core_properties.title=TITLE
doc.core_properties.author='Roberto Enrique Quezada-Rodríguez; Melanie Celeste Tello Fuentes'
doc.core_properties.subject='Semisynthetic TA-01 evaluation adapted for RGN Bulletin'
doc.core_properties.comments=''
FILE=OUT/'Manuscrito_TA01_RGN.docx'
doc.save(FILE)

# Restore all untouched package components byte-for-byte. The original media
# and opaque customXml parts remain intact; only the scoped parts are changed.
editable={'[Content_Types].xml','word/document.xml','word/_rels/document.xml.rels',
          'word/header1.xml','word/header2.xml','word/footer1.xml','word/footer2.xml',
          'word/settings.xml','docProps/core.xml','docProps/app.xml'}
with ZipFile(TEMPLATE) as z:original_parts={n:z.read(n) for n in z.namelist()}
with ZipFile(FILE) as z:final_parts={n:z.read(n) for n in z.namelist()}
for name,value in original_parts.items():
    if name not in editable:final_parts[name]=value
with ZipFile(FILE,'w',ZIP_DEFLATED) as z:
    for name,value in final_parts.items():z.writestr(name,value)

# Integrity gates: literal metric cells are preserved, no training was run.
for i in [2,3,5]:
    original=[[c.text for c in r.cells] for r in source.tables[i].rows]
    translated=TABLES[i]
    numeric_columns=[2,3,4,5] if i==2 else [0,2,3,4] if i==3 else [2,3,4]
    for a,b in zip(original[1:],translated[1:]):
        for col in numeric_columns:assert a[col]==b[col],(i,col,a[col],b[col])
for a,b in zip(source.tables[4].rows[1:],TABLES[4][1:]):assert a.cells[2].text==b[2]
assert len(ABSTRACT.split())<=250
checked=Document(FILE)
assert len(checked.tables)==7
assert len(checked.inline_shapes)==6
assert len(checked.element.xpath('//m:oMath'))==4
with ZipFile(FILE) as z:
    for name,value in original_parts.items():
        if name not in editable:assert z.read(name)==value,name
for n,info in PROVENANCE.items():
    for file,expected in info['sha256'].items():assert hashlib.sha256((REPO/file).read_bytes()).hexdigest()==expected
quality={'template_sha256':TEMPLATE_HASH,'source_unchanged':True,'metric_cells_unchanged':True,
         'abstract_words':len(ABSTRACT.split()),'croatian_abstract_words':len(CROATIAN_ABSTRACT.split()),
         'main_tables':6,'appendix_tables':1,'figure_count':6,'native_equations':4,
         'reference_count':46,'preserve_only_parts_identical':True,
         'figures':{x.name:{'pixels':Image.open(x).size,'dpi':Image.open(x).info.get('dpi')} for x in FIG.glob('*.jpg')}}
(OUT/'quality-check.json').write_text(json.dumps(quality,indent=2))
print(FILE)
print(json.dumps(quality,indent=2))
