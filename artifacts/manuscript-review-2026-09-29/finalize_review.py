from pathlib import Path
import json,re,zipfile,hashlib
from collections import defaultdict
from docx import Document
from lxml import etree
from PIL import Image
from openpyxl import load_workbook
OUT=Path(__file__).parent
original=json.loads((OUT/'original-extracted.json').read_text())
changes=json.loads((OUT/'changes.json').read_text())
d=Document(OUT/'Articulo_TA01_revisado.docx')
sections=defaultdict(list)
for c in changes:sections[c['section']].append(c)
parts=['# Cambios editoriales por sección','Copia de revisión, 29 de septiembre de 2026. El manuscrito completo está en Articulo_TA01_revisado.docx. Cambios resaltados en amarillo y comentarios de Word. Se excluyen de este registro cambios puramente tipográficos de cada celda o estilo, ya descritos en el informe. Los párrafos sin cambio textual conservan el contenido del original.']
for section,items in sections.items():
    parts.append('## '+section)
    for i,c in enumerate(items,1):
        parts += [f'### Cambio {i} · ubicación original: {c["paragraph"]}', '**Antes**', c['before'] or '(Párrafo vacío)', '**Después**',c['after'] or '(Párrafo vacío)','**Justificación:** '+c['reason']]
(OUT/'Cambios_por_seccion.md').write_text('\n\n'.join(parts)+'\n')
tabs=json.loads((OUT/'tables-export.json').read_text())
for t in tabs:
    for row in t['rows']:
        for i,v in enumerate(row):row[i]=v.replace('remuestreos por bloques de fecha','remuestreos por conglomerados de fecha')
(OUT/'tables-export.json').write_text(json.dumps(tabs,ensure_ascii=False,indent=2))
qa={}
origdoc=Document('/home/robertoa1/Descargas/Artículo - Prototipo semisintético FEM LSTM con corrección guiada por física para pronosticar desplazamientos en taludes.docx')
for oi,ni in [(2,2),(3,3)]:
    a=[[c.text for c in row.cells] for row in origdoc.tables[oi].rows]
    b=[[c.text for c in row.cells] for row in d.tables[ni].rows]
    assert a==b,('Tabla alterada',oi,ni)
qa['original_performance_tables_unchanged']=True
def nums(t):return re.findall(r'(?<!\w)\d+(?:[,.]\d+)?(?:\s*×\s*10[⁻⁰¹²³⁴⁵⁶⁷⁸⁹]+)?',t)
# Verificar cifras en resultados/conclusiones: excluye referencias cruzadas de Tabla.
for oi in [68,70,76,77,78,79,81,83,97,98,99]:
    before=origdoc.paragraphs[oi].text
    record=next((c for c in reversed(changes) if c['paragraph']==oi),None)
    after=record['after'] if record else before
    old=nums(re.sub(r'(?:Tabla|Figura)\s+\d+|TA-01','Referencia',before))
    new=nums(re.sub(r'(?:Tabla|Figura)\s+\d+|TA-01','Referencia',after))
    assert old==new,(oi,old,new)
qa['result_paragraph_numbers_unchanged']=True
qa['abstract_words']={lang:len(next(p.text for p in d.paragraphs if p.text.startswith(start)).split()) for lang,start in [('es','Introducción. La actualización'),('en','Introduction. Frequent')]}
assert all(v<=250 for v in qa['abstract_words'].values())
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(OUT/'Articulo_TA01_revisado.docx') as z:
    comments=etree.fromstring(z.read('word/comments.xml'));doc=etree.fromstring(z.read('word/document.xml'))
    ids=comments.xpath('//w:comment/@w:id',namespaces=ns)
    for tag in ['commentRangeStart','commentRangeEnd','commentReference']:
        assert sorted(doc.xpath('//w:'+tag+'/@w:id',namespaces=ns))==sorted(ids)
    qa['word_comments']=len(ids)
    qa['figure_images_and_equations']=len(d.inline_shapes)
qa['main_tables']=len(tabs);qa['total_tables_including_two_appendices']=len(d.tables)
qa['figure_resolution']={}
for f in sorted((OUT/'figuras').glob('*.png')):
    im=Image.open(f);dpi=im.info.get('dpi',(0,0));assert min(dpi)>=300
    qa['figure_resolution'][f.name]={'pixels':im.size,'dpi':dpi}
qa['manuscript_sha256']=hashlib.sha256((OUT/'Articulo_TA01_revisado.docx').read_bytes()).hexdigest()
(OUT/'quality-check.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
readme='Material gráfico TA-01 semisintético. Figura 1 es un esquema ilustrativo. Figuras 2–6 presentan artefactos archivados; no mediciones de mina. PNG a 400 dpi y SVG. Fuentes y hashes en figure-provenance.json. Las métricas no fueron recalculadas. Ver Informe_revision.md para límites y sustitución por datos instrumentales.\n'
with zipfile.ZipFile(OUT/'Material_figuras_y_trazabilidad.zip','w',zipfile.ZIP_DEFLATED) as z:
    z.writestr('LEEME.txt',readme)
    for f in sorted((OUT/'figuras').glob('*')):z.write(f,'figuras/'+f.name)
    for n in ['figure-provenance.json','hash-audit.json','quality-check.json','tables-export.json']:
        z.write(OUT/n,n)
print(json.dumps(qa,ensure_ascii=False))

