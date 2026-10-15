"""Produce a clean BGM-oriented article from the traced review copy."""

from __future__ import annotations

import importlib.util
import re
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Articulo_TA01_revisado.docx"
OUTPUT = ROOT / "Articulo_TA01_corregido_BGM.docx"
STRIP = Path("/home/robertoa1/.codex/plugins/cache/openai-primary-runtime/documents/26.909.12148/skills/documents/scripts/comments_strip.py")


def replace(paragraph, value: str) -> None:
    paragraph.text = value


def delete(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def prefix(paragraph, value: str) -> None:
    if not paragraph.text.startswith(value):
        paragraph.text = value + paragraph.text


def all_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


doc = Document(SOURCE)
p = list(doc.paragraphs)

# Correct the model equations to the code's standardized increment and clamp.
replace(
    p[50],
    "La LSTM recibe seis estados horarios y la lluvia futura conocida del experimento. "
    "El híbrido añade una corrección residual a la salida estandarizada de esa LSTM. "
    "En las ecuaciones (2) y (3), yₜ es el desplazamiento acumulado actual; h es el "
    "horizonte; qψ(Xₜ,ₕ) es el cambio estandarizado predicho por la LSTM; cφ(Zₜ,ₕ) "
    "es la corrección estandarizada; y μΔ y σΔ son la media y desviación del cambio "
    "objetivo, estimadas solo con el entrenamiento. X y Z reúnen las entradas "
    "temporales y del corrector, incluida la lluvia futura del experimento. El operador "
    "máximo impide que el pronóstico final disminuya respecto de yₜ."
)
for par, expr, number in (
    (p[51], "ŷLSTM(t+h) = max{yₜ, yₜ + μΔ + σΔ qψ(Xₜ,ₕ)}", "(2)"),
    (p[52], "ŷH(t+h) = max{yₜ, yₜ + μΔ + σΔ [qψ(Xₜ,ₕ) + cφ(Zₜ,ₕ)]}", "(3)"),
):
    replace(par, expr + "\t" + number)
    par.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT)
    for run in par.runs:
        run.font.name = "Cambria Math"
        run.font.size = Pt(11)

# Keep the archived bootstrap intervals but omit untraceable sign-flip p-values.
replace(
    p[66],
    "La diferencia pareada Δ = MAE(LSTM) − MAE(híbrido) se estimó mediante 5000 "
    "remuestreos por conglomerados de fecha de lluvia, con intervalos de confianza "
    "percentiles del 95 %. Agrupar por fecha conserva la dependencia de las ventanas "
    "de un mismo evento (Field & Welsh, 2007), aunque presupone que los conglomerados "
    "son suficientemente independientes entre fechas. No es un bootstrap de bloques "
    "de días consecutivos. No se reportan contrastes de hipótesis ni valores p porque "
    "no se conserva en el paquete experimental el cálculo reproducible de la prueba "
    "de inversión de signo previamente descrita. La prueba de robustez perturbó "
    "únicamente las entradas reservadas: ruido gaussiano del 5 % de la desviación "
    "de entrenamiento y pérdida aleatoria del 30 %, imputada con medias de "
    "entrenamiento. Las cifras proceden de predicciones guardadas del proyecto, "
    "no de observaciones de campo."
)
replace(
    p[89],
    "Al mover el año de prueba, el orden entre LSTM e híbrido cambió (Tabla 4). "
    "En 2024, el MAE del híbrido fue 1,15 % mayor a 1 h y 1,89 % mayor a 6 h; "
    "en 2025 fue 3,89 % menor a 1 h y 4,63 % menor a 6 h. Los cuatro intervalos "
    "del 95 % incluyeron cero. Estas estimaciones no sustentan una mejora consistente "
    "del corrector y tampoco prueban equivalencia de los modelos."
)
replace(
    p[90],
    "Los intervalos del 95 % para Δ fueron [−1,44 × 10⁻⁷; 4,25 × 10⁻⁸] mm a 1 h "
    "y [−1,15 × 10⁻⁶; 6,64 × 10⁻⁸] mm a 6 h en 2024; en 2025 fueron "
    "[−2,59 × 10⁻⁷; 6,88 × 10⁻⁷] mm y [−3,83 × 10⁻⁷; 6,25 × 10⁻⁶] mm. "
    "Las 82 fechas de prueba de 2024 y las 76 de 2025 constituyeron las unidades "
    "de remuestreo. Los intervalos describen variación entre fechas semisintéticas "
    "y no cuantifican incertidumbre geotécnica real."
)
replace(
    p[134],
    "Validez de conclusión. Las fechas, no las ventanas horarias, fueron las "
    "unidades de remuestreo. No se registró un cálculo de potencia ni un umbral de "
    "equivalencia antes del análisis; la inclusión de cero en los cuatro intervalos "
    "indica que la ventaja del híbrido no fue consistente entre cortes y horizontes, "
    "sin demostrar equivalencia. Los subconjuntos de lluvia ≥5 mm/día incluyen "
    "solo dos fechas en 2024 y tres en 2025; sus estimaciones de cobertura son "
    "inestables. La precisión de almacenamiento del FEM y la variación entre "
    "semillas de entrenamiento no se evaluaron como factores de sensibilidad."
)
replace(
    p[139],
    "RQ2 no encontró una ventaja consistente de la corrección residual: la LSTM fue "
    "mejor en 2024 y el híbrido en 2025, mientras que los cuatro intervalos pareados "
    "incluyeron cero. RQ3 mostró que ambos modelos se degradaron con ruido y faltantes, "
    "pero el híbrido perdió más MAE relativo y absoluto en las condiciones evaluadas. "
    "La calidad de entrada es, por tanto, una condición crítica del prototipo."
)
delete(p[172])  # Holm is no longer cited or claimed.

# Explicit archive mesh; do not present an incomplete command as a full recipe.
replace(
    p[69],
    "La reproducción parte del CSV diario NASA POWER conservado en data/rainfall/. "
    "Los 500 escenarios archivados se generaron con semilla 20260915 y malla "
    "30 × 20; esta última no coincide con la malla predeterminada del generador "
    "actual (14 × 9). Para replicar el experimento se deben fijar explícitamente "
    "--scenarios=500, --seed=20260915, --mesh-x=30 y --mesh-y=20 en "
    "npm run generate:fem --, y usar un directorio de salida nuevo para no "
    "sobrescribir los artefactos conservados. Después se crean las particiones "
    "por fecha y los cortes cronológicos documentados en docs/CASO_TA01_FEM.md; "
    "se entrenan ridge, LSTM y corrector por horizonte, se ejecutan las validaciones "
    "FEM y de robustez, y se consolida el informe. Los manifiestos guardan la "
    "procedencia y los parámetros de las corridas archivadas."
)
replace(
    p[70],
    "La recreación íntegra exige además reproducir los prefijos de salida, "
    "horizontes y pesos guardados de cada corte. Los scripts npm y el manifiesto "
    "de replicación son la referencia para esos argumentos; el comando de "
    "generación por sí solo no reproduce todas las tablas del estudio."
)

# Remove provisional editorial placeholders; authorship-specific assertions stay unassigned.
replace(
    p[144],
    "Ética. El estudio utiliza precipitación pública y datos geotécnicos "
    "simulados; no involucra participantes humanos, animales ni información personal."
)
delete(p[147])
replace(
    p[148],
    "Uso de inteligencia artificial generativa. OpenAI Codex asistió en el cotejo "
    "de esta versión con el código y los artefactos de sslope, así como en la revisión "
    "de redacción y formato. Los datos geotécnicos provienen del simulador y los "
    "modelos se entrenan mediante los scripts documentados. Los autores son "
    "responsables de verificar y aprobar las cifras, el método, las referencias "
    "y el texto final antes del envío."
)

# Three source records were checked against publisher pages / original paper.
replace(p[155], p[155].text.replace(" [VERIFICAR en fuente editorial]", ""))
replace(
    p[169],
    "Gómez Intriago, E. C., Guevara Sánchez, D. F., & Ortiz Hernández, E. "
    "(2024). Diseño de estabilidad de talud del sector de la Simbocal, vía "
    "Margarita-San Vicente, provincia de Manabí. Polo del Conocimiento, "
    "9(3), 1905–1917. https://doi.org/10.23857/pc.v9i3.6761"
)
replace(p[187], p[187].text.replace(" [VERIFICAR paginación]", ""))
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if "Guevara-Sánchez, D. F., Gómez-Intriago, E. C." in cell.text:
                cell.text = cell.text.replace(
                    "Guevara-Sánchez, D. F., Gómez-Intriago, E. C. y Ortiz-Hernández, E.",
                    "Gómez Intriago, E. C., Guevara Sánchez, D. F. y Ortiz Hernández, E."
                )

# Unverified verbatim quotations are not needed to support the analytical article.
delete(p[202])
delete(p[203])
last_table = doc.tables[-1]
last_table._element.getparent().remove(last_table._element)
replace(
    p[199],
    "Este material suplementario reúne la matriz de 40 antecedentes utilizada "
    "para contextualizar el trabajo. Se mantiene fuera del cuerpo principal "
    "para preservar el cierre de resultados, discusión y conclusiones."
)

# Running title and bilingual key points requested in current BGM instructions.
run_title = p[2].insert_paragraph_before("Título abreviado: Pronóstico semisintético FEM–LSTM en taludes")
run_title.style = doc.styles["Normal"]
before = p[12]
for line in (
    "Puntos clave",
    "• Se evalúa un prototipo FEM–LSTM sobre un talud y desplazamientos semisintéticos.",
    "• La ventaja del corrector cambia entre las pruebas cronológicas de 2024 y 2025.",
    "• La consistencia interna no equivale a validación geotécnica de campo.",
    "Key points",
    "• A FEM–LSTM prototype is evaluated on a semisynthetic slope and displacement series.",
    "• The correction's advantage changes between the 2024 and 2025 chronological tests.",
    "• Internal consistency does not constitute field geotechnical validation.",
):
    q = before.insert_paragraph_before(line)
    q.style = doc.styles["Normal"]

major = 0
minor = 0
for par in doc.paragraphs:
    if par.style.name == "Heading 1":
        if par.text in {"Referencias", "Material suplementario A", "Declaraciones"}:
            continue
        major += 1
        minor = 0
        prefix(par, f"{major}. ")
    elif par.style.name == "Heading 2" and major > 0:
        minor += 1
        prefix(par, f"{major}.{minor}. ")

for section in doc.sections:
    sectpr = section._sectPr
    existing = sectpr.find(qn("w:lnNumType"))
    if existing is None:
        existing = OxmlElement("w:lnNumType")
        sectpr.append(existing)
    existing.set(qn("w:countBy"), "1")
    existing.set(qn("w:start"), "1")
    existing.set(qn("w:restart"), "continuous")

for par in all_paragraphs(doc):
    for run in par.runs:
        run.font.highlight_color = None

spec = importlib.util.spec_from_file_location("comments_strip", STRIP)
module = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(module)
with tempfile.TemporaryDirectory(prefix="sslope-final-docx-") as temp:
    intermediate = Path(temp) / "edited.docx"
    doc.save(intermediate)
    stats = module.strip_comments(str(intermediate), str(OUTPUT))

print(OUTPUT)
print(stats)
