"""Build a journal-structured working copy without changing archived results."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


root = Path(__file__).resolve().parent
source = root / "Articulo_TA01_corregido_BGM.docx"
output = root / "Articulo_TA01_corregido_ESRJ.docx"
doc = Document(source)
p = list(doc.paragraphs)


def drop(par):
    par._element.getparent().remove(par._element)


# ESRJ requires English only. Keep the existing English title and abstract,
# and remove BGM-only bilingual front matter rather than duplicating it.
p[0].text = p[1].text.strip()
drop(p[1])
drop(p[2])
for index in [7, 8, 9, 13, 14, 15, 16, 17, 18, 19, 20]:
    drop(p[index])
p[10].text = "ABSTRACT"
p[21].text = "Abbreviations"
p[22].text = (
    "FEM: finite element method; FoS: factor of safety; CI: confidence interval; "
    "IoT: Internet of Things; LHS: Latin hypercube sampling; LSTM: long short-term "
    "memory; MAE: mean absolute error; MLP: multilayer perceptron; MSE: mean "
    "squared error; PIELM: physics-informed extreme learning machine; PINN: "
    "physics-informed neural network; RBF: radial basis function; RMSE: root "
    "mean squared error; RQ: research question."
)

# Put the source-data description ahead of the method, then separate sections.
for moved in p[50:53]:
    p[44]._p.addprevious(moved._p)
p[43].text = "DATA"
p[50].text = "TA-01 case and data provenance"
method = p[44].insert_paragraph_before("METHOD")
method.style = doc.styles["Heading 1"]

heading_map = {
    23: "INTRODUCTION",
    31: "Positioning against related work",
    36: "Research questions and contribution",
    44: "Prototype architecture and traceability",
    53: "Mechanical model and prediction target",
    56: "Mathematical formulation of the coupling",
    64: "Temporal models and residual correction",
    68: "Documented experimental configuration",
    73: "Validation and statistical analysis",
    77: "Reproducibility procedure",
    81: "RESULTS",
    82: "RQ4 Internal verification of the FEM",
    89: "RQ1 Predictive comparison by rainfall date",
    94: "RQ1 and RQ2 Chronological tests and uncertainty",
    112: "RQ4 Spatial approximation to discrete equilibrium",
    117: "RQ3 Robustness to input perturbations",
    133: "DISCUSSION",
    139: "Threats to validity",
    144: "Responsible use and governance",
    146: "CONCLUSIONS",
    150: "Data, code, and materials availability",
    152: "DECLARATIONS",
    158: "REFERENCES",
    205: "APPENDIX A",
}
for idx, text in heading_map.items():
    p[idx].text = text

# Statements about author ethics and AI are retained without invented CRediT roles.
for par in list(doc.paragraphs):
    if par.text.startswith("Ética."):
        par.text = (
            "Ethics. The study uses public precipitation data and simulated "
            "geotechnical data. It does not involve human participants, animals, "
            "or personal information."
        )
    elif par.text.startswith("Financiación."):
        par.text = "Funding. The authors report no specific funding, grants, or institutional support for this study."
    elif par.text.startswith("Conflictos de interés."):
        par.text = "Conflicts of interest. The authors report no conflicts of interest related to this work."
    elif par.text.startswith("Uso responsable."):
        par.text = (
            "Responsible use. This research prototype must not be used for safety, "
            "evacuation, or production decisions without independent validation "
            "and professional oversight."
        )

# The journal requires AI disclosure in Acknowledgements, not as an isolated
# declaration. The exact scope will be completed after the translation choice.
ai_par = next(q for q in doc.paragraphs if q.text.startswith("Uso de inteligencia artificial generativa."))
drop(ai_par)
ack = p[158].insert_paragraph_before("ACKNOWLEDGEMENTS")
ack.style = doc.styles["Heading 1"]
ack2 = p[158].insert_paragraph_before(
    "OpenAI Codex assisted with manuscript editing, consistency checks against "
    "the archived code and results, and document preparation. The authors retain "
    "responsibility for the final content."
)
ack2.style = doc.styles["Normal"]

# ESRJ does not request continuous line numbering.
for section in doc.sections:
    ln = section._sectPr.find(qn("w:lnNumType"))
    if ln is not None:
        section._sectPr.remove(ln)

doc.save(output)
print(output)
