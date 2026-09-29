from pathlib import Path
import json,re,math
from openpyxl import load_workbook
OUT=Path(__file__).parent
data=json.loads((OUT/'tables-export.json').read_text())
book=load_workbook(OUT/'outputs/manuscript-review-2026-09-29/Tablas_articulo_TA01.xlsx',data_only=False)
assert len(book.sheetnames)==6
count=0
for item in data:
    sh=book[f'Tabla {item["number"]}']
    for i,row in enumerate(item['rows']):
        for j,value in enumerate(row):
            expected=value.replace('remuestreos por bloques de fecha','remuestreos por conglomerados de fecha')
            actual=sh.cell(i+5,j+1).value
            if i and re.fullmatch(r'[-+]?\d+(?:[.,]\d+)?(?:e[-+]?\d+)?',expected,re.I):
                assert isinstance(actual,(int,float)) and math.isclose(float(expected.replace(',','.')),actual,rel_tol=1e-14,abs_tol=0),(item['number'],i,j,expected,actual)
            else:assert expected==actual,(item['number'],i,j,expected,actual)
            count+=1
    for row in sh:
        for cell in row:assert cell.data_type not in ['e','f'],(sh.title,cell.coordinate)
result={'sheets':book.sheetnames,'table_cells_matched':count,'excel_error_cells':0,'formulas':0,'purpose':'Copia editorial de métricas archivadas; no recalcular resultados.'}
(OUT/'excel-quality-check.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False))
