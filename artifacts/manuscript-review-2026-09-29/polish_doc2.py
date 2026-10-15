"""Final typography pass on the already clean article."""

from pathlib import Path
import re
import tempfile

from docx import Document
from docx.shared import Cm

path = Path(__file__).resolve().parent / "Articulo_TA01_corregido_BGM.docx"
doc = Document(path)
items = list(doc.paragraphs)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            items.extend(cell.paragraphs)

for p in items:
    for label in ("Palabras clave: ", "Keywords: "):
        if p.text.startswith(label):
            terms = p.text.removeprefix(label).rstrip(".").split("; ")
            p.text = label + "; ".join(term[:1].upper() + term[1:] for term in terms) + "."
    for run in p.runs:
        run.text = re.sub(r"(?<=\d),(?=\d)", ".", run.text)
    if p.style.name == "Normal" and p.text and not p.text.startswith(
        ("Figura ", "Figure ", "Tabla ", "Table ", "Nota.",
         "Palabras clave:", "Keywords:", "Título abreviado:", "• ")
    ):
        p.paragraph_format.first_line_indent = Cm(0.5)

with tempfile.TemporaryDirectory(prefix="sslope-polish-", dir=path.parent) as temp:
    intermediary = Path(temp) / "polished.docx"
    doc.save(intermediary)
    intermediary.replace(path)
print(path)
