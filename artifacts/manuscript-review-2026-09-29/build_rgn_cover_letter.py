"""Create the RGN editor letter and verified reviewer proposal separately from the paper."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "rgn-2026-09-30"
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.9)
for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
    style = doc.styles[name]
    style.font.name = "Times New Roman"
    fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
    for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
        fonts.set(qn("w:" + key), "Times New Roman")
    for key in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
        fonts.attrib.pop(qn("w:" + key), None)
    ppr = style.element.find(qn("w:pPr"))
    if ppr is not None:
        for border in list(ppr.findall(qn("w:pBdr"))):
            ppr.remove(border)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11 if name == "Normal" else 15 if name == "Title" else 12)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.08
    style.paragraph_format.widow_control = True

def p(text, style=None):
    par = doc.add_paragraph(text, style)
    for run in par.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.size = Pt(15 if style == "Title" else 12 if style == "Heading 1" else 11)
    props = par._p.get_or_add_pPr()
    for border in list(props.findall(qn("w:pBdr"))):
        props.remove(border)
    return par

def link(label, url):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(4)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), doc.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    props.append(fonts)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "22"); props.append(sz)
    run.append(props)
    t = OxmlElement("w:t"); t.text = label; run.append(t)
    h.append(run); par._p.append(h)
    return par

p("Editor letter and suggested reviewers", "Title")
p("30 September 2026")
p("Editor-in-Chief\nRudarsko-geološko-naftni zbornik")
p("Dear Editor-in-Chief,")
p('We request consideration of our manuscript, “Semisynthetic evaluation of FEM LSTM slope displacement forecasts”, by Roberto Enrique Quezada-Rodríguez and Melanie Celeste Tello Fuentes, as an original research article.')
p("The manuscript reports our original computational analyses of a controlled semisynthetic slope testbed. It combines a two-dimensional finite-element mechanical generator, long short-term memory forecasts and a physics-guided residual corrector. NASA POWER rainfall data provide environmental forcing; the displacements and assumed geotechnical conditions are simulated, not measurements from an operating mine.")
p("The contribution is a reproducible evaluation of forecasting performance, the added value of the corrector and sensitivity to input quality. The results do not establish consistent superiority of the corrector or validate field deployment. This scope is relevant to the journal’s coverage of computational geotechnics and engineering geology. The code and supporting validation files are available at https://github.com/RobertoA1/sslope.")
p("We received no funding for this research and declare no conflicts of interest. The manuscript discloses the use of OpenAI Codex. It is written in British English and has received AI-assisted language editing; professional or native-speaker proofreading has not yet been documented. We do not claim that AI-assisted editing satisfies the journal’s proofreading requirement.")
p("We propose the two specialists listed on the following page, selected for their complementary expertise and institutional affiliations outside our university. They have not been contacted. Their availability and appointment remain entirely at the discretion of the editors.")
p("Yours sincerely,\nRoberto Enrique Quezada-Rodríguez\nCorresponding author\nUniversidad Nacional de Trujillo\nCiudad Universitaria, Av. Juan Pablo II S/N, Urb. San Andrés\nTrujillo, La Libertad, Peru\nrequezadaro@unitru.edu.pe")

doc.add_page_break()
p("Suggested reviewers", "Title")
p("Proposal for the manuscript Semisynthetic evaluation of FEM LSTM slope displacement forecasts")

p("Jinsong Huang", "Heading 1")
p("Degree and position: Ph.D. in Civil Engineering; Professor, School of Engineering, University of Newcastle, Australia.")
p("Scientific fields: computational geotechnics, slope stability and geotechnical risk; machine learning and physics-informed modelling of unsaturated infiltration.")
p("Institutional email: Jinsong.Huang@newcastle.edu.au")
p("Suitability: his work on physics-informed learning with sparse and noisy observations is pertinent to the hybrid forecasting approach, uncertainty assessment and robustness tests. His research on rainfall infiltration and heterogeneous slopes also supports scrutiny of the physical assumptions.")
p("Selected work: co-author of “Reconstructing unsaturated infiltration behavior with sparse data via physics-informed deep learning” (Computers and Geotechnics, 168, 106162, 2024). Recent activity: co-author of “A rainfall infiltration-redistribution model for heterogeneous slopes considering multi-parameter uncertainties” (Journal of Rock Mechanics and Geotechnical Engineering, 17, 7882–7895, 2025).")
link("Institutional profile and contact", "https://profiles.newcastle.edu.au/Jinsong-Huang")
link("Physics-informed learning paper 10.1016/j.compgeo.2024.106162", "https://doi.org/10.1016/j.compgeo.2024.106162")
link("Recent paper 10.1016/j.jrmge.2025.04.005", "https://doi.org/10.1016/j.jrmge.2025.04.005")

p("Gianvito Scaringi", "Heading 1")
p("Degree and position: Ph.D.; Assistant Professor, Institute of Hydrogeology, Engineering Geology and Applied Geophysics, Faculty of Science, Charles University, Prague, Czechia.")
p("Scientific fields: engineering geology, geotechnical engineering, landslide processes, coupled behaviour of geomaterials and slope stability under environmental forcing.")
p("Institutional email: gianvito.scaringi@natur.cuni.cz")
p("Suitability: his experimental, monitoring and modelling research on landslides is pertinent to the realism of the semisynthetic scenarios, hydro-mechanical interpretation and distinction between numerical verification and field validation.")
p("Selected recent work: Loche, M., and Scaringi, G. (2025). “Assessing the influence of temperature on slope stability in a temperate climate: A nationwide spatial probability analysis in Italy”. Environmental Modelling & Software, 183, 106217.")
link("University profile", "https://cuni.cz/UKEN-2124.html")
link("Research group and affiliation", "https://geohazards.group/about/")
link("Institutional source for email and ORCID", "https://mdc.natur.cuni.cz/data/geohazards/pdf/DMPs/dmp_wp4.pdf")
link("Recent paper 10.1016/j.envsoft.2024.106217", "https://doi.org/10.1016/j.envsoft.2024.106217")

doc.core_properties.title = "RGN editor letter and suggested reviewers"
doc.core_properties.author = "Roberto Enrique Quezada-Rodríguez; Melanie Celeste Tello Fuentes"
doc.core_properties.subject = "TA-01 semisynthetic FEM LSTM manuscript"
OUT.mkdir(parents=True, exist_ok=True)
target = OUT / "Carta_editor_RGN_con_revisores.docx"
doc.save(target)
print(target)
